from datetime import datetime
from functools import wraps
from io import BytesIO
import base64
import os
import random
import secrets
import smtplib
import sqlite3
import string
import tempfile
import time
from email.message import EmailMessage

from flask import Flask, flash, redirect, render_template, request, send_file, session, url_for
from flask_sqlalchemy import SQLAlchemy
from sqlalchemy import func, text
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape, portrait
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import SimpleDocTemplate, Spacer, Table, TableStyle, Paragraph
from reportlab.pdfbase.pdfmetrics import stringWidth


app = Flask(__name__)
app.config["SECRET_KEY"] = "ticket-secret-key"
app.config["SQLALCHEMY_DATABASE_URI"] = "sqlite:///database.db"
app.config["SQLALCHEMY_TRACK_MODIFICATIONS"] = False
app.config["MAX_CONTENT_LENGTH"] = 50 * 1024 * 1024
app.config["SESSION_TIMEOUT_SECONDS"] = 120
db = SQLAlchemy(app)

TRANSLATIONS = {
    "fr": {
        "dashboard": "Tableau de bord", "events": "Événements", "categories": "Catégories",
        "users": "Utilisateurs", "sales": "Ventes", "tickets": "Billets", "control": "Contrôle",
        "reports": "Rapports",
        "account": "Compte", "logout": "Déconnexion", "login": "Connexion", "language": "Langue",
        "currency": "Devise", "french": "Français", "english": "Anglais", "save": "Enregistrer",
        "settings_saved": "Préférences enregistrées.", "role": "Rôle", "exchange_rate": "Taux de conversion",
        "exchange_help": "Nombre de FC pour 1 dollar américain.", "fc": "FC", "usd": "$",
        "event_exchange_rate": "Taux de conversion de l'événement", "event_exchange_help": "1 $ = combien de FC ?",
        "Performance globale": "Performance globale", "Bienvenue": "Bienvenue", "Nouvelle vente": "Nouvelle vente",
        "Événements": "Événements", "Ventes": "Ventes", "Billets": "Billets", "Revenu total": "Revenu total",
        "Événements à venir": "Événements à venir", "Prix moyen billet": "Prix moyen billet",
        "Taux de remplissage": "Taux de remplissage", "Billets validés": "Billets validés",
        "Programmation active": "Programmation active", "Par vente": "Par vente", "Capacité moyenne": "Capacité moyenne",
        "Contrôle d’entrée": "Contrôle d’entrée", "Transactions": "Transactions", "Ventes récentes": "Ventes récentes",
        "Activité": "Activité", "Derniers billets générés": "Derniers billets générés", "Performance": "Performance",
        "Points clés": "Points clés", "Aucun événement pour le moment.": "Aucun événement pour le moment.",
        "Liste des événements": "Liste des événements", "Ajouter un événement": "Ajouter un événement",
        "Catégories et prix": "Catégories et prix", "Ajouter une catégorie": "Ajouter une catégorie",
        "Vente de billets": "Vente de billets", "Sélectionner": "Sélectionner", "Valider la vente": "Valider la vente",
        "Billets vendus": "Billets vendus", "Contrôle d’entrée": "Contrôle d’entrée", "Valider l’entrée": "Valider l’entrée",
        "Rapports de ventes": "Rapports de ventes", "Rapport journalier": "Rapport journalier", "Rapport mensuel": "Rapport mensuel",
        "Filtrer": "Filtrer", "Télécharger PDF": "Télécharger PDF", "Télécharger Excel": "Télécharger Excel",
        "Chiffre d’affaires": "Chiffre d’affaires", "Nombre de ventes": "Nombre de ventes", "Billets vendus": "Billets vendus",
        "Accueil - Billetterie": "Accueil - Billetterie", "Connexion": "Connexion", "Événements": "Événements",
        "Catégories & Prix": "Catégories & Prix", "Configuration - StadeControl": "Configuration - StadeControl",
        "Contrôle d'entrée": "Contrôle d'entrée", "Nouvelle catégorie": "Nouvelle catégorie",
        "Nouvel événement": "Nouvel événement", "QR Code": "QR Code", "Utilisateurs": "Utilisateurs",
        "Gestion des utilisateurs": "Gestion des utilisateurs", "Créer un utilisateur": "Créer un utilisateur",
        "Liste des utilisateurs": "Liste des utilisateurs", "Tableau de bord": "Tableau de bord",
        "Ventes": "Ventes", "Billets": "Billets", "StadeControl · Planning": "StadeControl · Planning",
        "StadeControl · Tarification": "StadeControl · Tarification", "Ventes par événement": "Ventes par événement",
        "Billets vendus par événement": "Billets vendus par événement", "Statut des billets": "Statut des billets",
        "Revenus par catégorie": "Revenus par catégorie", "Utilisation des capacités": "Utilisation des capacités",
        "Créer un événement": "Créer un événement", "Ajouter une catégorie": "Ajouter une catégorie",
        "Vente de billets": "Vente de billets", "Billet": "Billet",
        "about": "À propos", "configuration": "Configuration", "system": "Système", "operations": "Opérations", "À propos - StadeControl": "À propos - StadeControl", "SYSTÈME": "SYSTÈME", "OPÉRATIONS": "OPÉRATIONS",
        "DÉCONNEXION": "DÉCONNEXION", "Contrôle": "Contrôle", "Accueil": "Accueil",
        "Référence": "Référence", "Client": "Client", "Email": "Email", "Catégorie": "Catégorie", "Prix": "Prix",
        "Statut": "Statut", "QR": "QR", "Voir QR": "Voir QR", "Validé": "Validé", "En attente": "En attente",
        "Aucun billet vendu.": "Aucun billet vendu.", "Événement": "Événement", "Date": "Date", "Lieu": "Lieu",
        "Capacité": "Capacité", "Nom de l'événement": "Nom de l'événement", "Nom de la catégorie": "Nom de la catégorie",
        "Quantité": "Quantité", "Quantité disponible": "Quantité disponible", "Nom d'utilisateur": "Nom d'utilisateur",
        "Mot de passe": "Mot de passe", "Administrateur": "Admin Système", "Auditeur": "Auditeur", "Vendeur": "Vendeur",
        "Photo de l'utilisateur": "Photo de l'utilisateur", "Activer la caméra": "Activer la caméra", "Capturer la photo": "Capturer la photo",
        "Créer": "Créer", "Enregistrer": "Enregistrer", "Type de rapport": "Type de rapport", "Mois": "Mois",
        "Organisez vos offres, vos tarifs et vos stocks par événement.": "Organisez vos offres, vos tarifs et vos stocks par événement.",
        "Accès refusé pour votre rôle.": "Accès refusé pour votre rôle.", "Identifiants incorrects.": "Identifiants incorrects.",
        "Événement ajouté avec succès.": "Événement ajouté avec succès.", "Catégorie ajoutée.": "Catégorie ajoutée.",
        "Billet introuvable. Vérifiez la référence.": "Billet introuvable. Vérifiez la référence.",
        "Ce billet a déjà été validé à l’entrée.": "Ce billet a déjà été validé à l’entrée.",
        "Total": "Total", "Aucune vente récente.": "Aucune vente récente.", "Aucun billet pour le moment.": "Aucun billet pour le moment.",
        "Validés": "Validés", "Revenu (€)": "Revenu (€)",
        "Stade Communal · Version 3.0.2": "Stade Communal · Version 3.0.2",
        "StadeControl est une application professionnelle de gestion de billetterie du Stade Communal.": "StadeControl est une application professionnelle de gestion de billetterie du Stade Communal.",
        "StadeControl centralise l'ensemble du cycle de gestion : création des événements, vente des billets, contrôle des entrées, gestion des utilisateurs et rapports d'activité.": "StadeControl centralise l'ensemble du cycle de gestion : création des événements, vente des billets, contrôle des entrées, gestion des utilisateurs et rapports d'activité.",
        "La plateforme a été pensée pour offrir aux équipes du Stade Communal une expérience fiable, claire et efficace, depuis la préparation d'un événement jusqu'au suivi de ses performances.": "La plateforme a été pensée pour offrir aux équipes du Stade Communal une expérience fiable, claire et efficace, depuis la préparation d'un événement jusqu'au suivi de ses performances.",
        "Développé par": "Développé par", "Version": "Version", "Édition": "Édition",
        "welcome_user": "Bienvenue, {username} !", "sale_registered": "Vente enregistrée : {quantity} billet(s) créés.",
        "entry_validated": "Entrée validée pour {buyer}.", "duplicate_user": "Ce nom d’utilisateur existe déjà.",
        "user_created": "Utilisateur {username} créé avec le rôle {role}.", "required_user_fields": "Merci de remplir le nom et le mot de passe.",
        "backup_title": "Sauvegarde et restauration", "download_backup": "Télécharger une sauvegarde",
        "restore_backup": "Restaurer une sauvegarde", "choose_backup": "Choisir un fichier .db",
        "restore_action": "Restaurer les données", "backup_warning": "La restauration remplacera les données actuelles.",
        "backup_missing": "Sélectionnez un fichier de sauvegarde.", "backup_type": "Le fichier doit être une base SQLite (.db, .sqlite ou .sqlite3).",
        "backup_success": "Sauvegarde restaurée avec succès.", "backup_invalid": "La sauvegarde est invalide ou n’a pas pu être restaurée.",
        "user_created_sent": "Utilisateur {username} créé. Le mot de passe a été envoyé à son adresse email.",
        "user_created_local": "Utilisateur {username} créé. Mot de passe temporaire : {password}",
        "password_reset_sent": "Le mot de passe de {username} a été réinitialisé et envoyé par email.",
        "password_reset_local": "Mot de passe réinitialisé pour {username}. Mot de passe temporaire : {password}",
        "active": "Actif", "inactive": "Inactif", "deactivate": "Désactiver", "activate": "Activer",
        "cancel_sale": "Annuler la vente", "cancelled": "Annulée", "confirm_cancel": "Confirmer l’annulation de cette vente ?",
        "phone": "Téléphone", "reset_password": "Réinitialiser", "reset_confirm": "Réinitialiser le mot de passe de cet utilisateur ?", "Action": "Action",
    },
    "en": {
        "dashboard": "Dashboard", "events": "Events", "categories": "Categories", "users": "Users",
        "sales": "Sales", "tickets": "Tickets", "control": "Check-in", "account": "Account",
        "reports": "Reports",
        "logout": "Log out", "login": "Log in", "language": "Language", "currency": "Currency",
        "french": "French", "english": "English", "save": "Save", "settings_saved": "Preferences saved.",
        "role": "Role", "exchange_rate": "Exchange rate", "exchange_help": "Number of FC for 1 US dollar.",
        "fc": "FC", "usd": "$", "event_exchange_rate": "Event exchange rate",
        "event_exchange_help": "1 $ = how many FC?",
        "Performance globale": "Overall performance", "Bienvenue": "Welcome", "Nouvelle vente": "New sale",
        "Événements": "Events", "Ventes": "Sales", "Billets": "Tickets", "Revenu total": "Total revenue",
        "Événements à venir": "Upcoming events", "Prix moyen billet": "Average ticket price",
        "Taux de remplissage": "Occupancy rate", "Billets validés": "Checked-in tickets",
        "Programmation active": "Active schedule", "Par vente": "Per sale", "Capacité moyenne": "Average capacity",
        "Contrôle d’entrée": "Entry control", "Transactions": "Transactions", "Ventes récentes": "Recent sales",
        "Activité": "Activity", "Derniers billets générés": "Latest tickets generated", "Performance": "Performance",
        "Points clés": "Key insights", "Aucun événement pour le moment.": "No events yet.",
        "Liste des événements": "Event list", "Ajouter un événement": "Add event",
        "Catégories et prix": "Categories and prices", "Ajouter une catégorie": "Add category",
        "Vente de billets": "Ticket sales", "Sélectionner": "Select", "Valider la vente": "Confirm sale",
        "Billets vendus": "Sold tickets", "Contrôle d’entrée": "Entry control", "Valider l’entrée": "Validate entry",
        "Rapports de ventes": "Sales reports", "Rapport journalier": "Daily report", "Rapport mensuel": "Monthly report",
        "Filtrer": "Filter", "Télécharger PDF": "Download PDF", "Télécharger Excel": "Download Excel",
        "Chiffre d’affaires": "Revenue", "Nombre de ventes": "Number of sales", "Billets vendus": "Tickets sold",
        "Accueil - Billetterie": "Home - Ticketing", "Connexion": "Log in", "Événements": "Events",
        "Catégories & Prix": "Categories & Prices", "Configuration - StadeControl": "Configuration - StadeControl",
        "Contrôle d'entrée": "Entry control", "Nouvelle catégorie": "New category", "Nouvel événement": "New event",
        "QR Code": "QR Code", "Utilisateurs": "Users", "Gestion des utilisateurs": "User management",
        "Créer un utilisateur": "Create a user", "Liste des utilisateurs": "User list", "Tableau de bord": "Dashboard",
        "Ventes": "Sales", "Billets": "Tickets", "StadeControl · Planning": "StadeControl · Planning",
        "StadeControl · Tarification": "StadeControl · Pricing", "Ventes par événement": "Sales by event",
        "Billets vendus par événement": "Tickets sold by event", "Statut des billets": "Ticket status",
        "Revenus par catégorie": "Revenue by category", "Utilisation des capacités": "Capacity usage",
        "Créer un événement": "Create an event", "Ajouter une catégorie": "Add a category",
        "Vente de billets": "Ticket sales", "Billet": "Ticket",
        "about": "About", "configuration": "Configuration", "system": "System", "operations": "Operations", "À propos - StadeControl": "About - StadeControl", "SYSTÈME": "SYSTEM", "OPÉRATIONS": "OPERATIONS",
        "DÉCONNEXION": "LOG OUT", "Contrôle": "Check-in", "Accueil": "Home",
        "Référence": "Reference", "Client": "Customer", "Email": "Email", "Catégorie": "Category", "Prix": "Price",
        "Statut": "Status", "QR": "QR", "Voir QR": "View QR", "Validé": "Validated", "En attente": "Pending",
        "Aucun billet vendu.": "No tickets sold.", "Événement": "Event", "Date": "Date", "Lieu": "Location",
        "Capacité": "Capacity", "Nom de l'événement": "Event name", "Nom de la catégorie": "Category name",
        "Quantité": "Quantity", "Quantité disponible": "Available quantity", "Nom d'utilisateur": "Username",
        "Mot de passe": "Password", "Administrateur": "System Admin", "Auditeur": "Auditor", "Vendeur": "Seller",
        "Photo de l'utilisateur": "User photo", "Activer la caméra": "Enable camera", "Capturer la photo": "Capture photo",
        "Créer": "Create", "Enregistrer": "Save", "Type de rapport": "Report type", "Mois": "Month",
        "Organisez vos offres, vos tarifs et vos stocks par événement.": "Organize offers, prices and stock by event.",
        "Accès refusé pour votre rôle.": "Access denied for your role.", "Identifiants incorrects.": "Incorrect credentials.",
        "Événement ajouté avec succès.": "Event added successfully.", "Catégorie ajoutée.": "Category added.",
        "Billet introuvable. Vérifiez la référence.": "Ticket not found. Check the reference.",
        "Ce billet a déjà été validé à l’entrée.": "This ticket has already been checked in.",
        "Total": "Total", "Aucune vente récente.": "No recent sales.", "Aucun billet pour le moment.": "No tickets yet.",
        "Validés": "Validated", "Revenu (€)": "Revenue ($)",
        "Stade Communal · Version 3.0.2": "Stade Communal · Version 3.0.2",
        "StadeControl est une application professionnelle de gestion de billetterie du Stade Communal.": "StadeControl is a professional ticketing management application for Stade Communal.",
        "StadeControl centralise l'ensemble du cycle de gestion : création des événements, vente des billets, contrôle des entrées, gestion des utilisateurs et rapports d'activité.": "StadeControl centralizes the entire management cycle: event creation, ticket sales, entry control, user management and activity reports.",
        "La plateforme a été pensée pour offrir aux équipes du Stade Communal une expérience fiable, claire et efficace, depuis la préparation d'un événement jusqu'au suivi de ses performances.": "The platform gives Stade Communal teams a reliable, clear and efficient experience, from event preparation to performance monitoring.",
        "Développé par": "Developed by", "Version": "Version", "Édition": "Edition",
        "welcome_user": "Welcome, {username}!", "sale_registered": "Sale recorded: {quantity} ticket(s) created.",
        "entry_validated": "Entry validated for {buyer}.", "duplicate_user": "This username already exists.",
        "user_created": "User {username} created with role {role}.", "required_user_fields": "Please fill in the username and password.",
        "backup_title": "Backup and restore", "download_backup": "Download a backup",
        "restore_backup": "Restore a backup", "choose_backup": "Choose a .db file",
        "restore_action": "Restore data", "backup_warning": "Restoring will replace the current data.",
        "backup_missing": "Select a backup file.", "backup_type": "The file must be a SQLite database (.db, .sqlite or .sqlite3).",
        "backup_success": "Backup restored successfully.", "backup_invalid": "The backup is invalid or could not be restored.",
        "user_created_sent": "User {username} created. The password was sent to their email address.",
        "user_created_local": "User {username} created. Temporary password: {password}",
        "password_reset_sent": "{username}'s password was reset and sent by email.",
        "password_reset_local": "Password reset for {username}. Temporary password: {password}",
        "active": "Active", "inactive": "Inactive", "deactivate": "Deactivate", "activate": "Activate",
        "cancel_sale": "Cancel sale", "cancelled": "Cancelled", "confirm_cancel": "Confirm cancellation of this sale?",
        "phone": "Phone", "reset_password": "Reset password", "reset_confirm": "Reset this user's password?", "Action": "Action",
    },
}


def current_language():
    return session.get("language", "fr") if session.get("language", "fr") in TRANSLATIONS else "fr"


def translate(key):
    return TRANSLATIONS[current_language()].get(key, key)


def current_currency():
    return session.get("currency", "FC") if session.get("currency", "FC") in ("FC", "USD") else "FC"


def convert_amount(amount, event=None):
    amount = float(amount or 0)
    if current_currency() == "USD":
        rate = float(getattr(event, "taux_conversion", 2800) or 2800) if event else 2800
        return amount / rate
    return amount


def money(amount, event=None):
    symbol = "$" if current_currency() == "USD" else "FC"
    return f"{convert_amount(amount, event):,.2f} {symbol}"


@app.context_processor
def inject_preferences():
    return {
        "_": translate,
        "money": money,
        "current_currency": current_currency(),
        "current_language": current_language(),
    }


@app.before_request
def enforce_session_timeout():
    if "user_id" not in session or request.endpoint in {"login", "static"}:
        return None
    now = time.time()
    last_activity = session.get("last_activity", now)
    if now - last_activity >= app.config["SESSION_TIMEOUT_SECONDS"]:
        session.clear()
        flash("Votre session a expiré après 2 minutes d’inactivité.", "warning")
        return redirect(url_for("login"))
    session["last_activity"] = now
    return None


class User(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    username = db.Column(db.String(80), unique=True, nullable=False)
    password = db.Column(db.String(200), nullable=False)
    role = db.Column(db.String(50), default="admin")
    language = db.Column(db.String(5), default="fr")
    currency = db.Column(db.String(5), default="FC")
    photo = db.Column(db.Text, nullable=True)
    active = db.Column(db.Boolean, nullable=False, default=True)
    email = db.Column(db.String(150), nullable=False, default="")
    telephone = db.Column(db.String(30), nullable=False, default="")


class Evenement(db.Model):
    __tablename__ = "evenement"
    id = db.Column(db.Integer, primary_key=True)
    nom = db.Column(db.String(150), nullable=False)
    date = db.Column(db.String(20), nullable=False)
    lieu = db.Column(db.String(150), nullable=False)
    capacite = db.Column(db.Integer, nullable=False)
    description = db.Column(db.Text, default="")
    taux_conversion = db.Column(db.Float, nullable=False, default=2800)
    latitude = db.Column(db.Float, nullable=True)
    longitude = db.Column(db.Float, nullable=True)
    categories = db.relationship("Categorie", backref="evenement", cascade="all, delete-orphan")
    tickets = db.relationship("Ticket", backref="evenement", cascade="all, delete-orphan")
    ventes = db.relationship("Vente", backref="evenement", cascade="all, delete-orphan")


class Categorie(db.Model):
    __tablename__ = "categorie"
    id = db.Column(db.Integer, primary_key=True)
    evenement_id = db.Column(db.Integer, db.ForeignKey("evenement.id"), nullable=False)
    nom = db.Column(db.String(100), nullable=False)
    prix = db.Column(db.Float, nullable=False, default=0.0)
    quantite = db.Column(db.Integer, nullable=False, default=0)
    tickets = db.relationship("Ticket", backref="categorie", cascade="all, delete-orphan")


class Vente(db.Model):
    __tablename__ = "vente"
    id = db.Column(db.Integer, primary_key=True)
    evenement_id = db.Column(db.Integer, db.ForeignKey("evenement.id"), nullable=False)
    acheteur = db.Column(db.String(150), nullable=False)
    email = db.Column(db.String(150), nullable=False)
    total = db.Column(db.Float, nullable=False, default=0.0)
    created_at = db.Column(db.DateTime, default=datetime.utcnow)
    annulee = db.Column(db.Boolean, nullable=False, default=False)
    vendeur_id = db.Column(db.Integer, db.ForeignKey("user.id"), nullable=True)
    vendeur = db.relationship("User", backref="ventes_enregistrees")
    tickets = db.relationship("Ticket", backref="vente", cascade="all, delete-orphan")


class Ticket(db.Model):
    __tablename__ = "ticket"
    id = db.Column(db.Integer, primary_key=True)
    reference = db.Column(db.String(30), unique=True, nullable=False)
    evenement_id = db.Column(db.Integer, db.ForeignKey("evenement.id"), nullable=False)
    categorie_id = db.Column(db.Integer, db.ForeignKey("categorie.id"), nullable=False)
    vente_id = db.Column(db.Integer, db.ForeignKey("vente.id"), nullable=False)
    acheteur = db.Column(db.String(150), nullable=False)
    email = db.Column(db.String(150), nullable=False)
    prix = db.Column(db.Float, nullable=False)
    checked_in = db.Column(db.Boolean, default=False)
    created_at = db.Column(db.DateTime, default=datetime.utcnow)


def generate_reference():
    return "TCK-" + "".join(random.choices(string.ascii_uppercase + string.digits, k=8))


def generate_temporary_password():
    return secrets.token_urlsafe(9)


def send_user_password(user, temporary_password, reset=False):
    smtp_host = os.getenv("STADECONTROL_SMTP_HOST")
    smtp_port = int(os.getenv("STADECONTROL_SMTP_PORT", "587"))
    smtp_user = os.getenv("STADECONTROL_SMTP_USER")
    smtp_password = os.getenv("STADECONTROL_SMTP_PASSWORD")
    sender = os.getenv("STADECONTROL_SMTP_FROM", smtp_user or "")
    if not smtp_host or not smtp_user or not smtp_password or not sender or not user.email:
        return False

    message = EmailMessage()
    message["Subject"] = "StadeControl - Réinitialisation de votre accès" if reset else "StadeControl - Votre accès"
    message["From"] = sender
    message["To"] = user.email
    message.set_content(
        f"Bonjour {user.username},\n\n"
        f"Votre mot de passe temporaire StadeControl est : {temporary_password}\n\n"
        "Connectez-vous puis demandez à un administrateur de le modifier.\n"
        "StadeControl - Stade Communal"
    )
    with smtplib.SMTP(smtp_host, smtp_port, timeout=15) as smtp:
        smtp.starttls()
        smtp.login(smtp_user, smtp_password)
        smtp.send_message(message)
    return True


def login_required(f):
    @wraps(f)
    def decorated_function(*args, **kwargs):
        if "user_id" not in session:
            return redirect(url_for("login"))
        return f(*args, **kwargs)

    return decorated_function


def role_required(*allowed_roles):
    allowed_roles = tuple(role.lower() for role in allowed_roles)

    def decorator(f):
        @wraps(f)
        def decorated_function(*args, **kwargs):
            if "user_id" not in session:
                return redirect(url_for("login"))

            user = User.query.get(session["user_id"])
            if user is None or not user.active:
                session.clear()
                return redirect(url_for("login"))

            role = (user.role or "").lower()
            if role == "admin":
                role = "administrateur"
                user.role = "administrateur"
                db.session.commit()

            if role not in allowed_roles:
                flash(translate("Accès refusé pour votre rôle."), "danger")
                return redirect(url_for("dashboard"))
            session["user_role"] = role
            return f(*args, **kwargs)

        return decorated_function

    return decorator


@app.route("/")
def index():
    if "user_id" in session:
        return redirect(url_for("dashboard"))
    return render_template("index.html")


@app.route("/accueil")
def accueil():
    if "user_id" in session:
        return redirect(url_for("dashboard"))
    return render_template("index.html")


@app.route("/login", methods=["GET", "POST"])
def login():
    if request.method == "POST":
        username = request.form.get("username", "").strip()
        password = request.form.get("password", "")
        user = User.query.filter_by(username=username, password=password).first()
        if user and not user.active:
            flash("Ce compte est désactivé. Contactez un administrateur.", "danger")
        elif user:
            if (user.role or "").lower() == "admin":
                user.role = "administrateur"
                db.session.commit()
            session["user_id"] = user.id
            session["username"] = user.username
            session["user_role"] = (user.role or "administrateur").lower()
            session["language"] = user.language or "fr"
            session["currency"] = user.currency or "FC"
            session["last_activity"] = time.time()
            flash(translate("welcome_user").format(username=user.username), "success")
            return redirect(url_for("dashboard"))
        flash(translate("Identifiants incorrects."), "danger")
    return render_template("login.html")


@app.route("/logout")
def logout():
    session.clear()
    return redirect(url_for("login"))


@app.route("/account", methods=["GET", "POST"])
@login_required
def account():
    user = User.query.get(session["user_id"])
    if request.method == "POST":
        user.language = request.form.get("language", "fr") if request.form.get("language") in TRANSLATIONS else "fr"
        user.currency = request.form.get("currency", "FC") if request.form.get("currency") in ("FC", "USD") else "FC"
        db.session.commit()
        session["language"] = user.language
        session["currency"] = user.currency
        flash(translate("settings_saved"), "success")
        return redirect(url_for("account"))
    return render_template("account.html", user=user)


@app.route("/a-propos")
def a_propos():
    return render_template("a_propos.html")


@app.route("/configuration")
@login_required
@role_required("administrateur")
def configuration():
    return render_template("configuration.html")


def database_path():
    database_name = db.engine.url.database
    if os.path.isabs(database_name):
        return database_name
    return os.path.join(app.instance_path, database_name)


@app.route("/backup", methods=["GET", "POST"])
@login_required
@role_required("administrateur")
def backup():
    if request.method == "POST":
        uploaded_file = request.files.get("backup_file")
        if not uploaded_file or not uploaded_file.filename:
            flash(translate("backup_missing"), "danger")
            return redirect(url_for("backup"))
        if not uploaded_file.filename.lower().endswith((".db", ".sqlite", ".sqlite3")):
            flash(translate("backup_type"), "danger")
            return redirect(url_for("backup"))

        temporary_path = None
        try:
            with tempfile.NamedTemporaryFile(suffix=".db", delete=False) as temporary_file:
                temporary_path = temporary_file.name
                uploaded_file.save(temporary_path)

            source = sqlite3.connect(temporary_path)
            integrity = source.execute("PRAGMA integrity_check").fetchone()[0]
            table_names = {
                row[0] for row in source.execute("SELECT name FROM sqlite_master WHERE type='table'")
            }
            required_tables = {"user", "evenement", "categorie", "vente", "ticket"}
            if integrity != "ok" or not required_tables.issubset(table_names):
                raise ValueError("invalid_backup")

            target_path = database_path()
            db.session.remove()
            db.engine.dispose()
            target = sqlite3.connect(target_path)
            source.backup(target)
            target.close()
            source.close()
            flash(translate("backup_success"), "success")
        except (sqlite3.Error, ValueError, OSError):
            flash(translate("backup_invalid"), "danger")
        finally:
            if temporary_path and os.path.exists(temporary_path):
                os.remove(temporary_path)
        return redirect(url_for("backup"))

    return render_template("backup.html")


@app.route("/backup/download")
@login_required
@role_required("administrateur")
def download_backup():
    path = database_path()
    filename = f"stadecontrol-backup-{datetime.utcnow().strftime('%Y%m%d-%H%M%S')}.db"
    return send_file(path, as_attachment=True, download_name=filename, mimetype="application/octet-stream")


@app.route("/dashboard")
@login_required
@role_required("administrateur", "auditeur", "vendeur")
def dashboard():
    total_evenements = Evenement.query.count()
    total_ventes = Vente.query.filter_by(annulee=False).count()
    total_billets = Ticket.query.count()
    billets_valides = Ticket.query.filter_by(checked_in=True).count()
    total_revenu = db.session.query(func.coalesce(func.sum(Vente.total), 0)).filter(Vente.annulee.is_(False)).scalar() or 0
    ventes_recents = Vente.query.filter_by(annulee=False).order_by(Vente.created_at.desc()).limit(5).all()
    avg_ticket_price = db.session.query(func.avg(Ticket.prix)).scalar() or 0
    upcoming_events = 0
    today = datetime.utcnow().strftime("%Y-%m-%d")
    for evenement in Evenement.query.all():
        if evenement.date and evenement.date >= today:
            upcoming_events += 1
    occupancy = 0
    if total_evenements:
        occupancy = round((sum((Ticket.query.filter_by(evenement_id=e.id).count() / (e.capacite or 1)) for e in Evenement.query.all()) / total_evenements) * 100, 1)
    checked_rate = round((billets_valides / total_billets) * 100, 1) if total_billets else 0

    recent_activity = Ticket.query.order_by(Ticket.created_at.desc()).limit(4).all()

    event_sales = (
        db.session.query(Evenement.nom, func.count(Ticket.id).label("tickets_vendus"))
        .outerjoin(Ticket, Ticket.evenement_id == Evenement.id)
        .group_by(Evenement.id, Evenement.nom)
        .order_by(func.count(Ticket.id).desc())
        .all()
    )
    event_labels = [item[0] for item in event_sales]
    event_values = [item[1] for item in event_sales]

    status_billets = {
        translate("Validés"): billets_valides,
        translate("En attente"): max(total_billets - billets_valides, 0),
    }

    category_revenue = (
        db.session.query(Categorie.nom, func.coalesce(func.sum(Ticket.prix), 0).label("revenu"))
        .outerjoin(Ticket, Ticket.categorie_id == Categorie.id)
        .group_by(Categorie.id, Categorie.nom)
        .order_by(func.coalesce(func.sum(Ticket.prix), 0).desc())
        .all()
    )
    category_labels = [item[0] for item in category_revenue]
    category_values = [float(item[1]) for item in category_revenue]

    event_load = []
    for evenement in Evenement.query.order_by(Evenement.date).all():
        vendus = Ticket.query.filter_by(evenement_id=evenement.id).count()
        capacite = evenement.capacite or 1
        taux = min(round((vendus / capacite) * 100), 100) if capacite else 0
        event_load.append({
            "nom": evenement.nom,
            "vendus": vendus,
            "capacite": evenement.capacite,
            "taux": taux,
        })

    return render_template(
        "dashboard.html",
        total_evenements=total_evenements,
        total_ventes=total_ventes,
        total_billets=total_billets,
        billets_valides=billets_valides,
        total_revenu=total_revenu,
        ventes_recents=ventes_recents,
        event_labels=event_labels,
        event_values=event_values,
        status_billets=status_billets,
        category_labels=category_labels,
        category_values=category_values,
        event_load=event_load,
        avg_ticket_price=avg_ticket_price,
        upcoming_events=upcoming_events,
        occupancy=occupancy,
        checked_rate=checked_rate,
        recent_activity=recent_activity,
    )


def report_period():
    period = request.args.get("period", "month")
    if period not in ("day", "month"):
        period = "month"
    now = datetime.utcnow()
    if period == "day":
        selected = request.args.get("date", now.strftime("%Y-%m-%d"))
        try:
            start = datetime.strptime(selected, "%Y-%m-%d")
        except ValueError:
            selected = now.strftime("%Y-%m-%d")
            start = datetime.strptime(selected, "%Y-%m-%d")
        end = start.replace(hour=23, minute=59, second=59)
        label = selected
    else:
        selected = request.args.get("month", now.strftime("%Y-%m"))
        try:
            start = datetime.strptime(selected, "%Y-%m")
        except ValueError:
            selected = now.strftime("%Y-%m")
            start = datetime.strptime(selected, "%Y-%m")
        end = datetime(start.year + (start.month == 12), 1 if start.month == 12 else start.month + 1, 1)
        label = selected
    return period, selected, start, end, label


def report_data():
    period, selected, start, end, label = report_period()
    current_user = User.query.get(session["user_id"])
    seller_id = request.args.get("seller_id", type=int)
    query = Vente.query.filter(
        Vente.created_at >= start,
        Vente.created_at < end,
        Vente.annulee.is_(False),
    )
    if current_user.role == "vendeur":
        seller_id = current_user.id
        query = query.filter(Vente.vendeur_id == current_user.id)
    elif current_user.role == "administrateur" and seller_id:
        query = query.filter(Vente.vendeur_id == seller_id)
    sales = query.order_by(Vente.created_at.desc()).all()
    total_revenue = sum(float(sale.total or 0) for sale in sales)
    total_tickets = sum(len(sale.tickets) for sale in sales)
    seller = User.query.get(seller_id) if seller_id else None
    return {
        "period": period,
        "selected": selected,
        "label": label,
        "sales": sales,
        "total_revenue": total_revenue,
        "total_tickets": total_tickets,
        "seller_id": seller_id,
        "seller_name": seller.username if seller else (current_user.username if current_user.role == "vendeur" else "Tous les vendeurs"),
        "event_names": sorted({sale.evenement.nom for sale in sales}),
        "exporter": current_user.username,
        "exported_at": datetime.now(),
    }


@app.route("/rapports")
@login_required
@role_required("administrateur", "auditeur", "vendeur")
def rapports():
    current_user = User.query.get(session["user_id"])
    sellers = User.query.filter(User.role == "vendeur", User.active.is_(True)).order_by(User.username).all()
    return render_template("rapports.html", report=report_data(), sellers=sellers, report_user=current_user)


@app.route("/rapports/export/excel")
@login_required
@role_required("administrateur", "auditeur", "vendeur")
def export_excel():
    report = report_data()
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Rapport ventes"
    sheet.append(["StadeControl - Stade Communal"])
    sheet.append(["Rapport", report["label"]])
    sheet.append(["Stade", "Stade Communal"])
    sheet.append(["Evenement(s)", ", ".join(report["event_names"]) or "Tous les événements"])
    sheet.append(["Vendeur filtre", report["seller_name"]])
    sheet.append(["Exporté par", report["exporter"]])
    sheet.append(["Date et heure d'export", report["exported_at"].strftime("%Y-%m-%d %H:%M:%S")])
    sheet.append([])
    sheet.append(["Date", "Client", "Email", "Evenement", "Vendeur", "Billets", "Total"])
    for sale in report["sales"]:
        sheet.append([
            sale.created_at.strftime("%Y-%m-%d %H:%M"), sale.acheteur, sale.email,
            sale.evenement.nom, sale.vendeur.username if sale.vendeur else "Ancienne vente", len(sale.tickets), sale.total,
        ])
    sheet.append([])
    sheet.append(["Totaux", "", "", "", "", report["total_tickets"], report["total_revenue"]])
    sheet["A1"].font = Font(bold=True, size=16, color="FFFFFF")
    for cell in sheet["1:1"]:
        cell.fill = PatternFill("solid", fgColor="1D1B2D")
    for cell in sheet["9:9"]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="8B5E3C")
    for column in sheet.columns:
        sheet.column_dimensions[column[0].column_letter].width = min(max(len(str(cell.value or "")) for cell in column) + 2, 34)
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    filename = f"rapport-{report['period']}-{report['selected']}.xlsx"
    return send_file(output, as_attachment=True, download_name=filename, mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


@app.route("/rapports/export/pdf")
@login_required
@role_required("administrateur", "auditeur", "vendeur")
def export_pdf():
    report = report_data()
    output = BytesIO()
    document = SimpleDocTemplate(output, pagesize=portrait(A4), rightMargin=28, leftMargin=28, topMargin=82, bottomMargin=45)
    styles = getSampleStyleSheet()
    table_style = styles["Normal"].clone("report-table")
    table_style.fontSize = 7
    table_style.leading = 8
    header_style = styles["Normal"].clone("report-table-header")
    header_style.fontSize = 7
    header_style.leading = 8
    header_style.textColor = colors.white
    header_style.fontName = "Helvetica-Bold"
    metadata = [
        ["Stade", "Stade Communal", "Période", report["label"]],
        ["Événement(s)", ", ".join(report["event_names"]) or "Tous les événements", "Vendeur", report["seller_name"]],
        ["Exporté par", report["exporter"], "Date et heure", report["exported_at"].strftime("%d/%m/%Y %H:%M:%S")],
    ]
    metadata_table = Table(metadata, colWidths=[70, 190, 70, 190])
    metadata_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#F1E7DF")),
        ("BACKGROUND", (2, 0), (2, -1), colors.HexColor("#F1E7DF")),
        ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"),
        ("FONTNAME", (2, 0), (2, -1), "Helvetica-Bold"),
        ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#D9CEC5")),
        ("PADDING", (0, 0), (-1, -1), 6),
    ]))
    content = [Paragraph(f"Rapport des ventes : {report['label']}", styles["Heading2"]), metadata_table, Spacer(1, 14)]
    rows = [[Paragraph(label, header_style) for label in ["Date", "Client", "Email", "Événement", "Vendeur", "Billets", "Total"]]]
    for sale in report["sales"]:
        rows.append([
            Paragraph(sale.created_at.strftime("%d/%m/%Y %H:%M"), table_style),
            Paragraph(sale.acheteur, table_style), Paragraph(sale.email, table_style),
            Paragraph(sale.evenement.nom, table_style),
            Paragraph(sale.vendeur.username if sale.vendeur else "Ancienne vente", table_style),
            Paragraph(str(len(sale.tickets)), table_style), Paragraph(money(sale.total, sale.evenement), table_style),
        ])
    rows.append([Paragraph("Totaux", table_style), "", "", "", "", Paragraph(str(report["total_tickets"]), table_style), Paragraph(money(report["total_revenue"]), table_style)])
    table = Table(rows, colWidths=[63, 70, 92, 104, 75, 45, 62], repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1D1B2D")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#F1E7DF")),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#D9CEC5")),
        ("PADDING", (0, 0), (-1, -1), 5),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -2), [colors.white, colors.HexColor("#F8F8F8")]),
    ]))
    content.append(table)
    content.extend([
        Spacer(1, 55),
        Paragraph("CERTIFICATION", styles["Heading3"]),
        Spacer(1, 6),
        Paragraph(f"Généré par : {report['exporter']}    |    Stade : Stade Communal", styles["Normal"]),
        Spacer(1, 10),
        Paragraph("Revu par : ________________________________    Date : ____________________", styles["Normal"]),
        Spacer(1, 8),
        Paragraph("Approuvé par : ______________________________    Date : ____________________", styles["Normal"]),
    ])
    def draw_pdf_header_footer(canvas, doc):
        canvas.saveState()
        width, height = portrait(A4)
        canvas.setFillColor(colors.HexColor("#176B57"))
        canvas.circle(42, height - 35, 17, fill=1, stroke=0)
        canvas.setFillColor(colors.white)
        canvas.setFont("Helvetica-Bold", 14)
        canvas.drawCentredString(42, height - 40, "S")
        canvas.setFillColor(colors.HexColor("#1D1B2D"))
        canvas.setFont("Helvetica-Bold", 13)
        canvas.drawString(68, height - 31, "StadeControl - Stade Communal")
        canvas.setFont("Helvetica", 8.5)
        canvas.drawString(68, height - 45, f"Événement(s): {', '.join(report['event_names']) or 'Tous les événements'}")
        canvas.drawRightString(width - 28, height - 31, f"Exporté par: {report['exporter']}")
        canvas.drawRightString(width - 28, height - 45, f"Export: {report['exported_at'].strftime('%d/%m/%Y %H:%M:%S')}")
        canvas.setStrokeColor(colors.HexColor("#176B57"))
        canvas.line(28, height - 58, width - 28, height - 58)
        canvas.setStrokeColor(colors.HexColor("#D9CEC5"))
        canvas.line(28, 32, width - 28, 32)
        canvas.setFillColor(colors.HexColor("#555555"))
        canvas.setFont("Helvetica", 8)
        canvas.drawString(28, 19, "StadeControl - Stade Communal")
        canvas.drawRightString(width - 28, 19, f"Date d'exportation: {report['exported_at'].strftime('%d/%m/%Y')} | Page {doc.page}")
        canvas.restoreState()

    document.build(content, onFirstPage=draw_pdf_header_footer, onLaterPages=draw_pdf_header_footer)
    output.seek(0)
    filename = f"rapport-{report['period']}-{report['selected']}.pdf"
    return send_file(output, as_attachment=True, download_name=filename, mimetype="application/pdf")


@app.route("/evenements")
@login_required
@role_required("administrateur", "auditeur")
def evenements():
    liste_evenements = Evenement.query.order_by(Evenement.date).all()
    event_rows = []
    for evenement in liste_evenements:
        vendus = len(evenement.tickets)
        capacite = evenement.capacite or 1
        event_rows.append({
            "evenement": evenement,
            "vendus": vendus,
            "taux": min(round((vendus / capacite) * 100), 100),
            "categories": len(evenement.categories),
            "latitude": evenement.latitude,
            "longitude": evenement.longitude,
        })
    map_events = [
        {
            "nom": row["evenement"].nom,
            "lieu": row["evenement"].lieu,
            "date": row["evenement"].date,
            "latitude": row["latitude"],
            "longitude": row["longitude"],
        }
        for row in event_rows if row["latitude"] is not None and row["longitude"] is not None
    ]
    return render_template("evenements.html", evenements=event_rows, map_events=map_events)


@app.route("/evenements/ajouter", methods=["GET", "POST"])
@login_required
@role_required("administrateur")
def ajouter_evenement():
    if request.method == "POST":
        nouveau_evenement = Evenement(
            nom=request.form["nom"],
            date=request.form["date"],
            lieu=request.form["lieu"],
            capacite=int(request.form["capacite"]),
            description=request.form.get("description", ""),
            taux_conversion=float(request.form.get("taux_conversion", 2800) or 2800),
            latitude=float(request.form["latitude"]) if request.form.get("latitude") else None,
            longitude=float(request.form["longitude"]) if request.form.get("longitude") else None,
        )
        db.session.add(nouveau_evenement)
        db.session.commit()
        flash(translate("Événement ajouté avec succès."), "success")
        return redirect(url_for("evenements"))
    return render_template("nouvel_evenement.html")


@app.route("/categories")
@login_required
@role_required("administrateur")
def categories():
    evenements = Evenement.query.order_by(Evenement.date).all()
    selected_event_id = request.args.get("evenement_id", type=int)
    if selected_event_id:
        categories = Categorie.query.filter_by(evenement_id=selected_event_id).order_by(Categorie.nom).all()
    else:
        categories = Categorie.query.order_by(Categorie.nom).all()
    category_rows = [
        {
            "categorie": categorie,
            "vendus": len(categorie.tickets),
            "restants": max(categorie.quantite - len(categorie.tickets), 0),
        }
        for categorie in categories
    ]
    return render_template(
        "categories.html",
        evenements=evenements,
        categories=category_rows,
        selected_event_id=selected_event_id,
    )


@app.route("/categories/ajouter", methods=["GET", "POST"])
@login_required
@role_required("administrateur")
def ajouter_categorie():
    if request.method == "POST":
        categorie = Categorie(
            evenement_id=int(request.form["evenement_id"]),
            nom=request.form["nom"],
            prix=float(request.form["prix"]),
            quantite=int(request.form["quantite"]),
        )
        db.session.add(categorie)
        db.session.commit()
        flash(translate("Catégorie ajoutée."), "success")
        return redirect(url_for("categories", evenement_id=categorie.evenement_id))
    evenements = Evenement.query.order_by(Evenement.date).all()
    return render_template("nouvel_categorie.html", evenements=evenements)


@app.route("/ventes", methods=["GET", "POST"])
@login_required
@role_required("administrateur", "vendeur")
def ventes():
    evenements = Evenement.query.order_by(Evenement.date).all()
    categories = []
    selected_event_id = request.args.get("evenement_id", type=int)

    if request.method == "POST":
        evenement_id = int(request.form["evenement_id"])
        categorie_id = int(request.form["categorie_id"])
        categorie = Categorie.query.get_or_404(categorie_id)
        quantite = int(request.form["quantite"])
        acheteur = request.form["acheteur"]
        email = request.form["email"]

        vente = Vente(
            evenement_id=evenement_id,
            acheteur=acheteur,
            email=email,
            total=categorie.prix * quantite,
            vendeur_id=session.get("user_id"),
        )
        db.session.add(vente)
        db.session.flush()

        for _ in range(quantite):
            ticket = Ticket(
                reference=generate_reference(),
                evenement_id=evenement_id,
                categorie_id=categorie_id,
                vente_id=vente.id,
                acheteur=acheteur,
                email=email,
                prix=categorie.prix,
            )
            db.session.add(ticket)

        db.session.commit()
        flash(translate("sale_registered").format(quantity=quantite), "success")
        return redirect(url_for("billets"))

    if selected_event_id:
        categories = Categorie.query.filter_by(evenement_id=selected_event_id).all()

    return render_template(
        "ventes.html",
        evenements=evenements,
        categories=categories,
        ventes_liste=Vente.query.order_by(Vente.created_at.desc()).limit(30).all(),
        selected_event_id=selected_event_id,
    )


@app.route("/ventes/<int:vente_id>/annuler", methods=["POST"])
@login_required
@role_required("administrateur")
def annuler_vente(vente_id):
    vente = Vente.query.get_or_404(vente_id)
    if vente.annulee:
        flash("Cette vente est déjà annulée.", "warning")
        return redirect(url_for("ventes"))
    for ticket in vente.tickets:
        db.session.delete(ticket)
    vente.annulee = True
    db.session.commit()
    flash(f"La vente #{vente.id} a été annulée.", "success")
    return redirect(url_for("ventes"))


@app.route("/billets")
@login_required
@role_required("administrateur", "auditeur", "vendeur")
def billets():
    tickets = Ticket.query.order_by(Ticket.created_at.desc()).all()
    return render_template("billets.html", tickets=tickets)


@app.route("/qr/<int:ticket_id>")
@login_required
@role_required("administrateur", "auditeur", "vendeur")
def qr_code(ticket_id):
    ticket = Ticket.query.get_or_404(ticket_id)
    qr_url = f"https://api.qrserver.com/v1/create-qr-code/?size=220x220&data={ticket.reference}"
    return render_template("qr_code.html", ticket=ticket, qr_url=qr_url)


@app.route("/controle-entree", methods=["GET", "POST"])
@login_required
@role_required("administrateur", "auditeur")
def controle_entree():
    ticket = None
    message = None
    status = "info"

    if request.method == "POST":
        reference = request.form.get("reference", "").strip()
        ticket = Ticket.query.filter_by(reference=reference).first()

        if ticket is None:
            message = translate("Billet introuvable. Vérifiez la référence.")
            status = "danger"
        elif ticket.checked_in:
            message = translate("Ce billet a déjà été validé à l’entrée.")
            status = "warning"
        else:
            ticket.checked_in = True
            db.session.commit()
            message = translate("entry_validated").format(buyer=ticket.acheteur)
            status = "success"

    return render_template("controle_entree.html", ticket=ticket, message=message, status=status)


@app.route("/users", methods=["GET", "POST"])
@login_required
@role_required("administrateur")
def users():
    if request.method == "POST":
        username = request.form.get("username", "").strip()
        password = request.form.get("password", "").strip() or generate_temporary_password()
        email = request.form.get("email", "").strip()
        telephone = request.form.get("telephone", "").strip()
        role = request.form.get("role", "vendeur")
        photo = request.form.get("photo", "")
        if photo and not photo.startswith("data:image/"):
            photo = ""

        if username and password and email and telephone:
            if User.query.filter_by(username=username).first():
                flash(translate("duplicate_user"), "danger")
            else:
                user = User(username=username, password=password, role=role, photo=photo or None, email=email, telephone=telephone)
                db.session.add(user)
                db.session.commit()
                try:
                    sent = send_user_password(user, password)
                except (OSError, smtplib.SMTPException):
                    sent = False
                message_key = "user_created_sent" if sent else "user_created_local"
                flash(translate(message_key).format(username=username, role=role, password=password), "success")
        else:
            flash("Veuillez remplir le nom, l’email, le téléphone et le mot de passe." if current_language() == "fr" else "Please fill in the username, email, phone and password.", "danger")

    users_list = User.query.order_by(User.username).all()
    return render_template("users.html", users=users_list)


@app.route("/users/<int:user_id>/edit", methods=["GET", "POST"])
@login_required
@role_required("administrateur")
def edit_user(user_id):
    user = User.query.get_or_404(user_id)
    if request.method == "POST":
        username = request.form.get("username", "").strip()
        email = request.form.get("email", "").strip()
        telephone = request.form.get("telephone", "").strip()
        role = request.form.get("role", user.role)
        password = request.form.get("password", "").strip()
        photo = request.form.get("photo", "").strip()

        duplicate = User.query.filter(User.username == username, User.id != user.id).first()
        if not username or not email or not telephone:
            flash("Veuillez remplir le nom, l’email et le téléphone.", "danger")
        elif duplicate:
            flash(translate("duplicate_user"), "danger")
        elif user.id == session.get("user_id") and role != "administrateur":
            flash("Vous ne pouvez pas retirer votre propre rôle administrateur.", "warning")
        else:
            user.username = username
            user.email = email
            user.telephone = telephone
            user.role = role
            if user.id != session.get("user_id"):
                user.active = request.form.get("active") == "on"
            if password:
                user.password = password
            if photo.startswith("data:image/"):
                user.photo = photo
            db.session.commit()
            if user.id == session.get("user_id"):
                session["username"] = user.username
                session["user_role"] = user.role
            flash("Utilisateur modifié avec succès.", "success")
            return redirect(url_for("users"))
    return render_template("edit_user.html", user=user)


@app.route("/users/<int:user_id>/reset-password", methods=["POST"])
@login_required
@role_required("administrateur")
def reset_user_password(user_id):
    user = User.query.get_or_404(user_id)
    if user.role not in ("auditeur", "vendeur"):
        flash("Seuls les comptes auditeur et vendeur peuvent être réinitialisés.", "warning")
        return redirect(url_for("users"))
    temporary_password = generate_temporary_password()
    user.password = temporary_password
    db.session.commit()
    try:
        sent = send_user_password(user, temporary_password, reset=True)
    except (OSError, smtplib.SMTPException):
        sent = False
    message_key = "password_reset_sent" if sent else "password_reset_local"
    flash(translate(message_key).format(username=user.username, password=temporary_password), "success")
    return redirect(url_for("users"))


@app.route("/users/<int:user_id>/toggle-active", methods=["POST"])
@login_required
@role_required("administrateur")
def toggle_user_active(user_id):
    user = User.query.get_or_404(user_id)
    if user.id == session.get("user_id"):
        flash("Vous ne pouvez pas désactiver votre propre compte.", "warning")
        return redirect(url_for("users"))
    user.active = not user.active
    db.session.commit()
    flash(f"Le compte {user.username} est maintenant {'actif' if user.active else 'inactif'}.", "success")
    return redirect(url_for("users"))


with app.app_context():
    db.create_all()

    evenement_columns = [col[1] for col in db.session.execute(text("PRAGMA table_info(evenement)"))]
    if "description" not in evenement_columns:
        db.session.execute(text("ALTER TABLE evenement ADD COLUMN description TEXT DEFAULT ''"))
        db.session.commit()
    if "taux_conversion" not in evenement_columns:
        db.session.execute(text("ALTER TABLE evenement ADD COLUMN taux_conversion FLOAT DEFAULT 2800"))
        db.session.commit()
    if "latitude" not in evenement_columns:
        db.session.execute(text("ALTER TABLE evenement ADD COLUMN latitude FLOAT"))
        db.session.commit()
    if "longitude" not in evenement_columns:
        db.session.execute(text("ALTER TABLE evenement ADD COLUMN longitude FLOAT"))
        db.session.commit()

    user_columns = [col[1] for col in db.session.execute(text("PRAGMA table_info(user)"))]
    if "role" not in user_columns:
        db.session.execute(text("ALTER TABLE user ADD COLUMN role VARCHAR(50) DEFAULT 'administrateur'"))
        db.session.commit()
    if "language" not in user_columns:
        db.session.execute(text("ALTER TABLE user ADD COLUMN language VARCHAR(5) DEFAULT 'fr'"))
        db.session.commit()
    if "currency" not in user_columns:
        db.session.execute(text("ALTER TABLE user ADD COLUMN currency VARCHAR(5) DEFAULT 'FC'"))
        db.session.commit()
    if "photo" not in user_columns:
        db.session.execute(text("ALTER TABLE user ADD COLUMN photo TEXT"))
        db.session.commit()
    if "active" not in user_columns:
        db.session.execute(text("ALTER TABLE user ADD COLUMN active BOOLEAN DEFAULT 1"))
        db.session.commit()
    if "email" not in user_columns:
        db.session.execute(text("ALTER TABLE user ADD COLUMN email VARCHAR(150) DEFAULT ''"))
        db.session.commit()
    if "telephone" not in user_columns:
        db.session.execute(text("ALTER TABLE user ADD COLUMN telephone VARCHAR(30) DEFAULT ''"))
        db.session.commit()

    vente_columns = [col[1] for col in db.session.execute(text("PRAGMA table_info(vente)"))]
    if "annulee" not in vente_columns:
        db.session.execute(text("ALTER TABLE vente ADD COLUMN annulee BOOLEAN DEFAULT 0"))
        db.session.commit()
    if "vendeur_id" not in vente_columns:
        db.session.execute(text("ALTER TABLE vente ADD COLUMN vendeur_id INTEGER"))
        db.session.commit()

    for user in User.query.all():
        if (user.role or "").lower() == "admin":
            user.role = "administrateur"
        if not user.language:
            user.language = "fr"
        if not user.currency:
            user.currency = "FC"
    for evenement in Evenement.query.all():
        if not evenement.taux_conversion:
            evenement.taux_conversion = 2800
    db.session.commit()

    if not User.query.filter_by(username="admin").first():
        db.session.add(User(username="admin", password="admin123", role="administrateur"))
        db.session.commit()


if __name__ == "__main__":
    app.run(debug=True)
